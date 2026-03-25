"""
Maintenance Card AI Parser — parse PDF/Word maintenance cards via GPT-5.4.

Flow: upload file → extract text → GPT parse → ParsedMaintenanceCard → confirm → save to DB.
Logs every attempt to CardParseLog.
"""
from __future__ import annotations

import asyncio
import base64
import json
import logging
import os
import time
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openai
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from config import settings
from models.base import async_session
from models.maintenance_lifecycle import (
    CardInterval,
    CardParseLog,
    CardSparePart,
    CardWorkItem,
    EquipmentCardLink,
    MaintenanceCardV2,
)
from schemas.maintenance import ParsedMaintenanceCard, ParsedEquipmentInfo, ParsedInterval

logger = logging.getLogger("scada.maintenance")

# ── GPT prompt ────────────────────────────────────────────────────────

PARSE_MAINTENANCE_CARD_PROMPT = """\
You are an expert maintenance engineer. Parse the following maintenance card document \
and extract ALL maintenance intervals, work items, and spare parts.

Return a valid JSON object with this EXACT structure:
{
  "equipment_info": {
    "manufacturer": "string or null",
    "equipment_type": "string or null — e.g. generator, pump, compressor",
    "model": "string or null",
    "description": "string or null"
  },
  "intervals": [
    {
      "name": "string — human-readable name, e.g. ТО-1, ТО-2, Ежедневное",
      "code": "string — short code: TO-1, TO-2, ETO, KR, etc.",
      "interval_value": "integer or null — hours between maintenance",
      "interval_type": "string — one of: hours, calendar_days, hours_or_days, once",
      "calendar_days": "integer or null — days between maintenance (for calendar/combo)",
      "labor_hours": "float or null — estimated labor hours",
      "is_periodic": true,
      "is_overhaul": false,
      "includes": ["array of codes this interval includes, e.g. TO-1 includes ETO"],
      "specific_work_items": ["array of work description strings SPECIFIC to this interval only"],
      "spare_parts": [
        {
          "part_number": "string or null",
          "name": "string",
          "unit": "string, default шт.",
          "quantity": 1.0,
          "model_filter": "string or null — if part is model-specific"
        }
      ]
    }
  ],
  "notes": "string or null — any additional notes from the document"
}

RULES:
1. Extract ALL intervals mentioned in the document (ETO, TO-1, TO-2, ..., KR/overhaul).
2. For each interval, list ONLY its OWN specific work items — do NOT duplicate inherited work.
3. If an interval includes another (e.g. TO-2 includes TO-1), put the included code in "includes".
4. Mark the overhaul (KR/капремонт) with "is_overhaul": true.
5. "once" type = first-time or commissioning maintenance.
6. If the document mentions both hours AND calendar days for one interval, use "hours_or_days".
7. Spare parts must be associated with their specific interval.
8. Use Russian names as they appear in the document.
9. Return ONLY the JSON, no markdown fences, no explanation.

DOCUMENT TEXT:
---
{document_text}
---
"""

# ── Text extraction helpers ───────────────────────────────────────────


def _extract_text_pdf(file_path: str) -> str:
    """Extract text from PDF using pdfplumber."""
    try:
        import pdfplumber

        texts: list[str] = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    texts.append(page_text)
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        if row:
                            texts.append("\t".join(str(c or "") for c in row))
        result = "\n".join(texts)
        # If pdfplumber extracted very little, try pymupdf as fallback
        if len(result) < 500:
            logger.info("pdfplumber extracted only %d chars, trying pymupdf fallback", len(result))
            try:
                import fitz  # pymupdf
                doc = fitz.open(file_path)
                mu_texts = []
                for page in doc:
                    mu_texts.append(page.get_text())
                doc.close()
                mu_result = "\n".join(mu_texts)
                if len(mu_result) > len(result):
                    logger.info("pymupdf extracted %d chars (vs pdfplumber %d)", len(mu_result), len(result))
                    return mu_result
            except Exception as mu_err:
                logger.warning("pymupdf fallback failed: %s", mu_err)
        return result
    except Exception as exc:
        logger.warning("pdfplumber extraction failed: %s, trying pymupdf", exc)
        try:
            import fitz
            doc = fitz.open(file_path)
            texts = [page.get_text() for page in doc]
            doc.close()
            return "\n".join(texts)
        except Exception:
            return ""


def _extract_text_docx(file_path: str) -> str:
    """Extract text from .docx using python-docx."""
    try:
        import docx

        doc = docx.Document(file_path)
        texts: list[str] = []
        for para in doc.paragraphs:
            if para.text.strip():
                texts.append(para.text)
        # Also extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = "\t".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    texts.append(row_text)
        return "\n".join(texts)
    except Exception as exc:
        logger.warning("python-docx extraction failed: %s", exc)
        return ""


async def _extract_via_vision(file_path: str) -> str:
    """Fallback: convert PDF pages to images and send to GPT-5.4 Vision."""
    try:
        from pdf2image import convert_from_path

        images = convert_from_path(file_path, dpi=200, first_page=1, last_page=10)
    except Exception as exc:
        logger.error("pdf2image conversion failed: %s", exc)
        return ""

    image_messages = []
    for i, img in enumerate(images):
        import io
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        b64 = base64.b64encode(buf.getvalue()).decode("utf-8")
        image_messages.append({
            "type": "image_url",
            "image_url": {"url": f"data:image/png;base64,{b64}", "detail": "high"},
        })

    client = openai.AsyncOpenAI(api_key=settings.OPENAI_API_KEY)
    try:
        resp = await client.chat.completions.create(
            model=settings.OPENAI_MODEL,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "text", "text": (
                        "Extract ALL text from these maintenance card pages. "
                        "Include tables as tab-separated rows. Return raw text only."
                    )},
                    *image_messages,
                ],
            }],
            max_completion_tokens=8000,
            temperature=0,
        )
        return resp.choices[0].message.content or ""
    except Exception as exc:
        logger.error("Vision text extraction failed: %s", exc)
        return ""


# ── Core parsing ──────────────────────────────────────────────────────


async def parse_card(file_path: str, method: str = "auto") -> dict:
    """Parse PDF/Word maintenance card.

    Returns:
        {"parsed": ParsedMaintenanceCard, "method": str, "confidence": float,
         "input_tokens": int, "output_tokens": int}
    """
    started_at = time.monotonic()
    ext = Path(file_path).suffix.lower()
    file_size = os.path.getsize(file_path) if os.path.exists(file_path) else 0

    # Step 1: extract text
    actual_method = method
    text = ""

    if ext == ".docx":
        text = _extract_text_docx(file_path)
        actual_method = "docx"
    elif ext in (".pdf", ".PDF"):
        if method in ("auto", "pdfplumber"):
            text = _extract_text_pdf(file_path)
            actual_method = "pdfplumber"
        if len(text) < 200 and method in ("auto", "vision"):
            logger.info("Text too short (%d chars), falling back to Vision", len(text))
            text = await _extract_via_vision(file_path)
            actual_method = "vision"
    else:
        # Try reading as plain text
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
            actual_method = "text"
        except Exception:
            raise ValueError(f"Unsupported file type: {ext}")

    if not text.strip():
        raise ValueError(f"Could not extract text from {file_path}")

    # Step 2: send to GPT-5.4 (timeout 240s — large cards generate 16k+ tokens)
    client = openai.AsyncOpenAI(api_key=settings.OPENAI_API_KEY, timeout=240.0)
    prompt = PARSE_MAINTENANCE_CARD_PROMPT.replace("{document_text}", text[:30000])

    parse_log = CardParseLog(
        file_name=Path(file_path).name,
        file_size_bytes=file_size,
        parse_method=actual_method,
        llm_model=settings.OPENAI_MODEL,
        started_at=datetime.utcnow(),
        success=False,
    )

    try:
        resp = await client.chat.completions.create(
            model=settings.OPENAI_MODEL,
            messages=[
                {"role": "system", "content": "You are a maintenance card parser. Return only valid JSON."},
                {"role": "user", "content": prompt},
            ],
            max_completion_tokens=16000,
            temperature=0,
            response_format={"type": "json_object"},
        )

        raw_content = resp.choices[0].message.content or "{}"
        # Check for truncated response (finish_reason != "stop")
        finish_reason = resp.choices[0].finish_reason if resp.choices else "unknown"
        if finish_reason == "length":
            logger.warning("GPT response truncated (max_tokens reached), output may be incomplete")
        input_tokens = resp.usage.prompt_tokens if resp.usage else 0
        output_tokens = resp.usage.completion_tokens if resp.usage else 0

        # Parse JSON response (tolerant — GPT may return slightly different structure)
        # Strip markdown code fences if present
        clean = raw_content.strip()
        if clean.startswith("```"):
            clean = clean.split("\n", 1)[-1]
            if clean.endswith("```"):
                clean = clean[:-3].strip()
        try:
            raw_json = json.loads(clean)
        except json.JSONDecodeError as jde:
            logger.warning("JSON parse failed (truncated?): %s. Attempting repair...", str(jde)[:100])
            # Try to repair truncated JSON by closing brackets
            repair = clean
            open_braces = repair.count("{") - repair.count("}")
            open_brackets = repair.count("[") - repair.count("]")
            repair += "]" * max(0, open_brackets) + "}" * max(0, open_braces)
            try:
                raw_json = json.loads(repair)
                logger.info("JSON repaired successfully")
            except json.JSONDecodeError:
                logger.error("JSON repair failed, returning empty result")
                raw_json = {"equipment_info": {}, "intervals": [], "notes": "JSON parse error — response truncated"}
        try:
            parsed = ParsedMaintenanceCard.model_validate(raw_json)
        except Exception as val_err:
            logger.warning("Pydantic validation failed, trying relaxed parse: %s", val_err)
            # Fallback: create with defaults
            parsed = ParsedMaintenanceCard(
                equipment_info=ParsedEquipmentInfo(**(raw_json.get("equipment_info") or {})),
                intervals=[ParsedInterval(**iv) for iv in (raw_json.get("intervals") or [])],
                notes=raw_json.get("notes"),
            )

        # Confidence heuristic
        confidence = _calculate_confidence(parsed, actual_method)

        # Auto-retry with Vision if no intervals found and not already using Vision
        # Timeout: 90s max for entire Vision retry (pdf2image + GPT call)
        if len(parsed.intervals) == 0 and actual_method != "vision" and ext in (".pdf", ".PDF"):
            logger.info(
                "No intervals found (confidence=%.2f, method=%s), retrying with Vision (90s timeout)",
                confidence, actual_method,
            )
            try:
                vr = await asyncio.wait_for(
                    _vision_retry(file_path, client, len(parsed.intervals), confidence),
                    timeout=90,
                )
                if vr and vr["intervals_count"] > len(parsed.intervals):
                    logger.info("Vision retry improved: intervals %d→%d, confidence %.2f→%.2f",
                                len(parsed.intervals), vr["intervals_count"], confidence, vr["confidence"])
                    parsed = vr["parsed"]
                    confidence = vr["confidence"]
                    actual_method = "vision_retry"
                    raw_json = vr["raw_json"]
                    input_tokens += vr.get("input_tokens", 0)
                    output_tokens += vr.get("output_tokens", 0)
            except asyncio.TimeoutError:
                logger.warning("Vision retry timed out (90s), using pdfplumber result")
            except Exception as vision_err:
                logger.warning("Vision retry failed: %s", vision_err)

        duration = time.monotonic() - started_at
        parse_log.success = True
        parse_log.input_tokens = input_tokens
        parse_log.output_tokens = output_tokens
        parse_log.total_cost_usd = Decimal(str(round(input_tokens * 0.00001 + output_tokens * 0.00003, 4)))
        parse_log.completed_at = datetime.utcnow()
        parse_log.duration_seconds = Decimal(str(round(duration, 2)))
        parse_log.raw_response = raw_json

        # Save parse log
        async with async_session() as session:
            session.add(parse_log)
            await session.commit()

        logger.info(
            "Card parsed: %s, method=%s, intervals=%d, confidence=%.2f",
            Path(file_path).name, actual_method, len(parsed.intervals), confidence,
        )

        return {
            "parsed": parsed,
            "method": actual_method,
            "confidence": confidence,
            "input_tokens": input_tokens,
            "output_tokens": output_tokens,
            "parse_log_id": parse_log.id,
        }

    except Exception as exc:
        duration = time.monotonic() - started_at
        parse_log.error_message = str(exc)[:1000]
        parse_log.completed_at = datetime.utcnow()
        parse_log.duration_seconds = Decimal(str(round(duration, 2)))

        async with async_session() as session:
            session.add(parse_log)
            await session.commit()

        logger.error("Card parsing failed: %s", exc)
        raise


async def _vision_retry(file_path: str, client, prev_count: int, prev_conf: float) -> dict | None:
    """Try Vision extraction + GPT parse. Returns dict with results or None."""
    vision_text = await _extract_via_vision(file_path)
    if not vision_text or len(vision_text) < 200:
        logger.info("Vision extraction returned too little text (%d chars)", len(vision_text or ""))
        return None

    prompt = PARSE_MAINTENANCE_CARD_PROMPT.replace("{document_text}", vision_text[:30000])
    resp = await client.chat.completions.create(
        model=settings.OPENAI_MODEL,
        messages=[
            {"role": "system", "content": "You are a maintenance card parser. Return only valid JSON."},
            {"role": "user", "content": prompt},
        ],
        max_completion_tokens=8000,
        temperature=0,
        response_format={"type": "json_object"},
    )
    content = resp.choices[0].message.content or "{}"
    clean = content.strip()
    if clean.startswith("```"):
        clean = clean.split("\n", 1)[-1]
        if clean.endswith("```"):
            clean = clean[:-3].strip()
    raw_json = json.loads(clean)
    parsed = ParsedMaintenanceCard.model_validate(raw_json)
    conf = _calculate_confidence(parsed, "vision")

    return {
        "parsed": parsed,
        "confidence": conf,
        "raw_json": raw_json,
        "intervals_count": len(parsed.intervals),
        "input_tokens": resp.usage.prompt_tokens if resp.usage else 0,
        "output_tokens": resp.usage.completion_tokens if resp.usage else 0,
    }


def _calculate_confidence(parsed: ParsedMaintenanceCard, method: str) -> float:
    """Heuristic confidence score 0.0-1.0."""
    score = 0.5
    if len(parsed.intervals) >= 2:
        score += 0.15
    if len(parsed.intervals) >= 4:
        score += 0.1
    if any(iv.spare_parts for iv in parsed.intervals):
        score += 0.1
    if any(iv.specific_work_items for iv in parsed.intervals):
        score += 0.1
    if parsed.equipment_info.manufacturer:
        score += 0.05
    if method == "pdfplumber":
        score += 0.05
    elif method == "vision":
        score -= 0.05
    return min(round(score, 2), 1.0)


# ── Save & link ───────────────────────────────────────────────────────


async def save_confirmed_card(
    parsed: dict, confirmed_by: str, session: AsyncSession
) -> MaintenanceCardV2:
    """Save confirmed parsed card to DB.

    Creates: MaintenanceCardV2 + CardInterval + CardWorkItem + CardSparePart.
    """
    data: ParsedMaintenanceCard = parsed["parsed"]
    eq = data.equipment_info

    card = MaintenanceCardV2(
        name=f"{eq.manufacturer or 'Unknown'} — {eq.equipment_type or 'equipment'}",
        manufacturer=eq.manufacturer,
        equipment_type=eq.equipment_type,
        model_filter=eq.model,
        status="confirmed",
        confirmed_by=confirmed_by,
        confirmed_at=datetime.utcnow(),
        parsed_at=datetime.utcnow(),
        parsed_by="gpt",
        parse_confidence=Decimal(str(parsed.get("confidence", 0.0))),
        raw_parsed_json=data.model_dump(),
        notes=data.notes,
    )
    session.add(card)
    await session.flush()  # get card.id

    for sort_idx, iv in enumerate(data.intervals):
        interval = CardInterval(
            card_id=card.id,
            name=iv.name,
            code=iv.code,
            interval_type=iv.interval_type,
            interval_hours=iv.interval_value,
            interval_days=iv.calendar_days,
            is_periodic=iv.is_periodic,
            is_overhaul=iv.is_overhaul,
            labor_hours=Decimal(str(iv.labor_hours)) if iv.labor_hours else None,
            includes=iv.includes or [],
            warn_threshold=None,
            task_threshold=None,
            sort_order=sort_idx,
        )
        session.add(interval)
        await session.flush()

        # Work items
        for wi_idx, work_desc in enumerate(iv.specific_work_items):
            session.add(CardWorkItem(
                interval_id=interval.id,
                description=work_desc,
                is_specific=True,
                sort_order=wi_idx,
            ))

        # Spare parts
        for sp_idx, sp in enumerate(iv.spare_parts):
            session.add(CardSparePart(
                interval_id=interval.id,
                part_number=sp.part_number,
                name=sp.name,
                unit=sp.unit,
                quantity=Decimal(str(sp.quantity)),
                model_filter=sp.model_filter,
                sort_order=sp_idx,
            ))

    # Update parse log with card_id
    if parsed.get("parse_log_id"):
        stmt = select(CardParseLog).where(CardParseLog.id == parsed["parse_log_id"])
        log_row = (await session.execute(stmt)).scalar_one_or_none()
        if log_row:
            log_row.card_id = card.id

    await session.commit()
    logger.info("Card saved: id=%d, name=%s, intervals=%d", card.id, card.name, len(data.intervals))
    return card


async def link_card_to_equipment(
    card_id: int, equipment_ids: list[int], linked_by: str, session: AsyncSession
) -> list[EquipmentCardLink]:
    """Link card to equipment units. Auto-transition: confirmed -> active."""
    links = []
    for eq_id in equipment_ids:
        link = EquipmentCardLink(
            equipment_id=eq_id,
            card_id=card_id,
            linked_by=linked_by,
        )
        session.add(link)
        links.append(link)

    # Auto-transition card status
    stmt = select(MaintenanceCardV2).where(MaintenanceCardV2.id == card_id)
    card = (await session.execute(stmt)).scalar_one_or_none()
    if card and card.status == "confirmed":
        card.status = "active"
        logger.info("Card %d auto-transitioned to 'active' on equipment link", card_id)

    await session.commit()
    logger.info("Card %d linked to %d equipment units by %s", card_id, len(equipment_ids), linked_by)
    return links


async def reparse_card(card_id: int, method: str, session: AsyncSession) -> dict:
    """Re-parse existing card. Status: active -> draft."""
    stmt = select(MaintenanceCardV2).where(MaintenanceCardV2.id == card_id)
    card = (await session.execute(stmt)).scalar_one_or_none()
    if not card:
        raise ValueError(f"Card {card_id} not found")

    if not card.source_file_path:
        raise ValueError(f"Card {card_id} has no source file path for re-parsing")

    # Reset status
    old_status = card.status
    card.status = "draft"
    await session.commit()
    logger.info("Card %d status changed %s -> draft for re-parsing", card_id, old_status)

    # Re-parse
    result = await parse_card(card.source_file_path, method=method)

    # Update parse log with card_id
    async with async_session() as log_session:
        if result.get("parse_log_id"):
            log_stmt = select(CardParseLog).where(CardParseLog.id == result["parse_log_id"])
            log_row = (await log_session.execute(log_stmt)).scalar_one_or_none()
            if log_row:
                log_row.card_id = card_id
                await log_session.commit()

    return result
